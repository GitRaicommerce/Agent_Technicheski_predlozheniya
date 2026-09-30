"""WP-08 / K-14 acceptance (T-16): the final DOCX obeys the calendar-date policy."""

from __future__ import annotations

import io
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock

import pytest
from docx import Document

from app.agents.proposal_timing import find_concrete_calendar_dates
from app.export.docx_generator import (
    CalendarDatesInExportError,
    generate_docx,
    scan_document_for_calendar_dates,
)


def _scalar(value):
    result = MagicMock()
    result.scalar_one_or_none.return_value = value
    return result


DATED_SCHEDULE = {
    "tasks": [
        {"uid": 1, "name": "Проектиране", "duration_days": 20.0, "start": "2026-10-06", "finish": "2026-10-25", "predecessors": ""},
        {"uid": 2, "name": "Строителство", "duration_value": 3.0, "duration_unit": "weeks", "start": "2026-10-26", "finish": "2026-11-16", "predecessors": "1"},
    ]
}


@pytest.mark.asyncio
async def test_docx_keeps_activities_and_durations_without_calendar_dates():
    project = SimpleNamespace(
        id="p1", name="Водопровод", contracting_authority=None, location=None, tender_date=None
    )
    db = MagicMock()
    db.get = AsyncMock(return_value=project)
    db.execute = AsyncMock(
        side_effect=[_scalar(None), _scalar(SimpleNamespace(schedule_json=DATED_SCHEDULE))]
    )

    content = await generate_docx("p1", db)

    doc = Document(io.BytesIO(content))
    all_text = "\n".join(p.text for p in doc.paragraphs)
    cells = [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    assert not find_concrete_calendar_dates(all_text)
    assert not any(find_concrete_calendar_dates(cell) for cell in cells)
    assert "Проектиране" in cells and "Строителство" in cells
    assert "20 дни" in cells and "3 седмици" in cells
    assert "1" in cells  # dependency kept
    assert "Начало" not in cells and "Край" not in cells


def test_whole_output_scan_finds_dates_in_text_and_tables():
    doc = Document()
    doc.add_paragraph("Корица: 2026-10-01")
    doc.add_paragraph("Изпълнението започва на 06.10.2026 г.")
    table = doc.add_table(rows=1, cols=1)
    table.rows[0].cells[0].text = "Край 25.10.2026"

    findings = scan_document_for_calendar_dates(doc, skip_paragraphs=1)

    locations = [finding["location"] for finding in findings]
    assert "paragraph:1" in locations
    assert any(location.startswith("table:0") for location in locations)
    assert "paragraph:0" not in locations  # cover metadata is excluded


@pytest.mark.asyncio
async def test_export_route_hard_blocks_a_dated_document(client, mock_db):
    from unittest.mock import patch

    from tests.conftest import _make_project

    project = _make_project()
    mock_db.get = AsyncMock(return_value=project)
    with (
        patch("app.routers.export._require_export_ready", new=AsyncMock(return_value={})),
        patch(
            "app.export.docx_generator.generate_docx",
            new=AsyncMock(side_effect=CalendarDatesInExportError([{"location": "paragraph:4", "dates": ["06.10.2026"]}])),
        ),
    ):
        response = await client.get(f"/api/v1/export/{project.id}/docx?allow_incomplete=true")

    assert response.status_code == 409
    detail = response.json()["detail"]
    assert detail["code"] == "concrete_calendar_dates"
    assert detail["can_export_current_draft"] is False
    assert detail["calendar_date_locations"][0]["location"] == "paragraph:4"
